#!/usr/bin/env python3
"""把 bench_preset_ab.py 产出的 preset_ab_report.json 汇总成一个 Excel 工作簿。

设计目标：一次给出「不同并发档位」的横向对照，便于回答
「CPU 比是否随并发漂移」这个核心问题。

工作表：
  * 汇总        —— 每份报告一行：并发/参数/几何平均墙钟比·CPU 比/核数/环境
  * 逐素材      —— 每份报告 × 每条素材：墙钟、CPU、比值、核数
  * 并发对比    —— 同一素材在不同并发档下的比值并排（关键判据）
  * 环境        —— 各报告的机器指纹（CPU/AVX2/ffmpeg/负载）

用法：
  python3 make_ab_xlsx.py --out preset_ab_report.xlsx a.json b.json ...
"""
import argparse
import json
import os

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def load(path):
    with open(path) as f:
        return json.load(f)


def label_of(rep):
    env = rep.get('environment', {})
    return 'c%s' % env.get('concurrency', '?')


def summary_rows(reports):
    rows = []
    for path, rep in reports:
        env = rep['environment']
        s = rep['summary']
        p = rep['params']
        rows.append({
            '档位': label_of(rep),
            '并发': env.get('concurrency'),
            '编码器': p['codec'],
            'A预设': p['preset_a'],
            'B预设': p['preset_b'],
            '锚点crf': p['anchor'],
            '每格计时次数': p['runs'],
            '预热次数': p['warmup'],
            '素材数': rep['n_materials'],
            '墙钟比_几何平均': s.get('wall_ratio_geomean'),
            'CPU比_几何平均': s.get('cpu_ratio_geomean'),
            'CPU比_中位': s.get('cpu_ratio_median'),
            'CPU比_最小': s.get('cpu_ratio_min'),
            'CPU比_最大': s.get('cpu_ratio_max'),
            'A侧核数_中位': s.get('cores_a_median'),
            'B侧核数_中位': s.get('cores_b_median'),
            '开测负载1min': (env.get('load_at_start') or [None])[0],
            '结束负载1min': (env.get('load_at_end') or [None])[0],
        })
    return rows


def per_material_rows(reports):
    rows = []
    for _, rep in reports:
        tag = label_of(rep)
        for m in rep['per_material']:
            rows.append({
                '档位': tag,
                '素材': m['material'],
                'A墙钟s': m['wall_a_s'],
                'B墙钟s': m['wall_b_s'],
                '墙钟比': m['wall_ratio'],
                'A_CPU_s': m['cpu_a_s'],
                'B_CPU_s': m['cpu_b_s'],
                'CPU比': m['cpu_ratio'],
                'A核数': m['cores_a'],
                'B核数': m['cores_b'],
                '墙钟波动s': m['wall_spread'],
            })
    return rows


def compare_rows(reports):
    """同一素材在不同档位下的比值并排。

    不同报告里素材路径前缀不同（全量跑是 `10s/xxx.mp4`，单条标定是
    `m/xxx.mp4`），故按**文件名**归并，才能跨档位对齐。
    """
    per = {}
    for _, rep in reports:
        tag = label_of(rep)
        for m in rep['per_material']:
            key = os.path.basename(m['material'])
            per.setdefault(key, {})[tag] = m
    tags = [label_of(r) for _, r in reports]
    rows = []
    for mat, by_tag in per.items():
        if len(by_tag) < 2:
            continue
        row = {'素材': mat}
        for t in tags:
            if t in by_tag:
                row['墙钟比@%s' % t] = by_tag[t]['wall_ratio']
                row['CPU比@%s' % t] = by_tag[t]['cpu_ratio']
        base = tags[0]
        if len(tags) > 1 and base in by_tag and tags[-1] in by_tag:
            row['CPU比漂移%%(%s→%s)' % (base, tags[-1])] = (
                (by_tag[tags[-1]]['cpu_ratio'] / by_tag[base]['cpu_ratio'] - 1) * 100)
            row['墙钟比漂移%%(%s→%s)' % (base, tags[-1])] = (
                (by_tag[tags[-1]]['wall_ratio'] / by_tag[base]['wall_ratio'] - 1) * 100)
        rows.append(row)
    return rows


def env_rows(reports):
    rows = []
    for _, rep in reports:
        env = rep['environment']
        rows.append({
            '档位': label_of(rep),
            'CPU型号': env.get('cpu_model'),
            '核数': env.get('cores'),
            'cpuinfo_avx2': env.get('cpuinfo_avx2'),
            'cpuinfo_avx512f': env.get('cpuinfo_avx512f'),
            'x265实际asm': env.get('x265_asm'),
            'x265用AVX2': env.get('x265_uses_avx2'),
            'ffmpeg': env.get('ffmpeg'),
            '开测负载': env.get('load_at_start'),
            '结束负载': env.get('load_at_end'),
        })
    return rows


def avx2_measure_rows(rep):
    """AVX2 反证实验的逐格测量表。"""
    arms = rep.get('arms', {})
    caps = rep.get('capabilities', {})
    res = rep.get('results', {})
    delta = rep.get('delta_cpu_pct', {}) or {}
    presets = sorted({k.split('|', 1)[1] for k in res})
    rows = []
    for arm in arms:
        for p in presets:
            c = res.get('%s|%s' % (arm, p))
            if not c:
                continue
            mask = arms[arm]
            rows.append({
                '臂': arm,
                '预设': p,
                '掩码': 'auto' if mask is None else mask,
                'CPU中位s': c.get('cpu_median'),
                '墙钟中位s': c.get('wall_median'),
                '相对自动ΔCPU%': (delta.get(p) if arm != 'auto' else 0.0),
                'x265能力': caps.get(arm),
            })
    return rows


def avx2_summary_rows(rep):
    caps = rep.get('capabilities', {})
    d = rep.get('delta_cpu_pct', {}) or {}
    return [
        {'项': '素材', '值': rep.get('clip')},
        {'项': 'crf', '值': rep.get('crf')},
        {'项': '每格计时次数', '值': rep.get('runs')},
        {'项': '自动臂能力(auto)', '值': caps.get('auto')},
        {'项': '去AVX2臂能力', '值': caps.get('noavx2')},
        {'项': '绝对代价 slow(%)', '值': d.get('slow')},
        {'项': '绝对代价 slower(%)', '值': d.get('slower')},
        {'项': '预设比值_自动', '值': rep.get('ratio_auto')},
        {'项': '预设比值_去AVX2', '值': rep.get('ratio_noavx2')},
        {'项': '比值漂移(%)', '值': rep.get('ratio_drift_pct')},
    ]


def write_sheet(writer, name, rows):
    df = pd.DataFrame(rows)
    df.to_excel(writer, sheet_name=name, index=False)
    return df


def style(path):
    from openpyxl import load_workbook
    wb = load_workbook(path)
    head = Font(bold=True, color='FFFFFF')
    fill = PatternFill('solid', fgColor='4472C4')
    for ws in wb.worksheets:
        for c in ws[1]:
            c.font = head
            c.fill = fill
            c.alignment = Alignment(horizontal='center', vertical='center')
        ws.freeze_panes = 'A2'
        for col in ws.columns:
            letter = get_column_letter(col[0].column)
            width = max(len(str(c.value)) if c.value is not None else 0
                        for c in col) + 4
            ws.column_dimensions[letter].width = min(max(width, 8), 42)
    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--out', required=True)
    ap.add_argument('--avx2', default=None,
                    help='可选的 avx2_falsify_report.json，追加 AVX2 反证工作表')
    ap.add_argument('reports', nargs='+')
    args = ap.parse_args()

    reports = [(p, load(p)) for p in args.reports]

    with pd.ExcelWriter(args.out, engine='openpyxl') as w:
        write_sheet(w, '汇总', summary_rows(reports))
        write_sheet(w, '逐素材', per_material_rows(reports))
        cmp = compare_rows(reports)
        if cmp:
            write_sheet(w, '并发对比', cmp)
        write_sheet(w, '环境', env_rows(reports))
        if args.avx2:
            a = load(args.avx2)
            write_sheet(w, 'AVX2反证', avx2_measure_rows(a))
            write_sheet(w, 'AVX2结论', avx2_summary_rows(a))

    style(args.out)
    print('已写：%s' % args.out)
    print('  汇总 %d 行 / 逐素材 %d 行 / 并发对比 %d 行'
          % (len(summary_rows(reports)), len(per_material_rows(reports)),
             len(compare_rows(reports))))


if __name__ == '__main__':
    main()
